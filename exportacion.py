"""
Módulo de Exportación de Resultados y Gráficos para Avance Físico de Obra.

Genera reportes corporativos en Excel (.xlsx) con estilos avanzados (openpyxl),
archivos CSV y empaquetado en lote de gráficos en archivo ZIP.
"""

import io
import zipfile
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, Any


def aplicar_estilo_general_excel(ws, titulo: str, idx_header: int = 3, tipo_tabla: str = "semanal"):
    """
    Aplica estilos corporativos premium a la hoja de cálculo de Excel:
    - Cabecera azul corporativo (#1F4E78) con texto blanco en negrita
    - Bordes sutiles y cuadrícula activa
    - Zebra striping suave en las filas de datos
    - Formateo numérico automático para porcentajes y enteros
    - Ajuste automático de ancho de columnas y congelación de paneles
    """
    ws.views.sheetView[0].showGridLines = True

    font_titulo = Font(name="Segoe UI", size=15, bold=True, color="1F4E78")
    font_header = Font(name="Segoe UI", size=10.5, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

    fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    border_side = Side(style="thin", color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # Título en A1
    ws.cell(row=1, column=1, value=titulo).font = font_titulo
    ws.row_dimensions[1].height = 26

    # Cabecera
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=idx_header, column=col_idx)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[idx_header].height = 26

    # Filas de datos
    for r_idx in range(idx_header + 1, ws.max_row + 1):
        if ws.cell(row=r_idx, column=1).value is None and ws.cell(row=r_idx, column=2).value is None:
            continue

        row_fill = fill_even if (r_idx % 2 == 0) else fill_odd
        ws.row_dimensions[r_idx].height = 20

        for c_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.fill = row_fill
            cell.border = cell_border
            cell.font = Font(name="Segoe UI", size=9.5)
            cell.alignment = Alignment(horizontal="left", vertical="center")

            val = cell.value

            if tipo_tabla == "semanal":
                if c_idx == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.number_format = '0'
                elif c_idx in [2, 3, 4, 5, 6]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, (int, float)):
                        cell.value = val / 100.0
                        cell.number_format = '0.00%'
                elif c_idx == 7:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    val_str = str(val).upper()
                    if "ADELANTADA" in val_str:
                        cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="006100")
                    elif "RETRASADA" in val_str:
                        cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="9C0006")
                    else:
                        cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="9C6500")

            elif tipo_tabla in ["especialidad", "avance_encargado"]:
                if c_idx == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx in [2, 3, 4]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, (int, float)):
                        cell.value = val / 100.0
                        cell.number_format = '0.00%'

            elif tipo_tabla == "partidas_encargado":
                if c_idx == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx == 2:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.number_format = '#,##0'

            elif tipo_tabla == "validaciones":
                if c_idx in [1, 2, 3]:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif c_idx == 4:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    val_str = str(val).upper()
                    if "PASÓ" in val_str or "OK" in val_str or "PASO" in val_str:
                        cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="006100")
                    else:
                        cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="9C0006")

    # Ajuste automático del ancho de columnas
    for col in ws.columns:
        max_len = 0
        for cell in col:
            if cell.row == 1:
                continue
            val_str = ""
            if cell.value is not None:
                if cell.number_format == '0.00%':
                    val_str = f"{cell.value * 100:.2f}%"
                else:
                    val_str = str(cell.value)
            max_len = max(max_len, len(val_str))
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    ws.freeze_panes = f"A{idx_header + 1}"


def generar_excel_completo(
    df_semanal: pd.DataFrame,
    df_especialidad: pd.DataFrame,
    df_partidas_encargado: pd.DataFrame,
    df_avance_encargado: pd.DataFrame,
    validaciones_dict: Dict[str, Any]
) -> bytes:
    """
    Genera un archivo Excel completo con 5 hojas corporativas estructuradas.
    """
    wb = openpyxl.Workbook()

    # 1. Avance_Semanal
    ws_as = wb.active
    ws_as.title = "Avance_Semanal"
    headers_as = [
        "Semana", "Avance Planificado Semanal", "Avance Real Semanal",
        "Avance Planificado Acumulado", "Avance Real Acumulado",
        "Diferencia Acumulada", "Estado de Avance"
    ]
    for c_idx, name in enumerate(headers_as, start=1):
        ws_as.cell(row=3, column=c_idx, value=name)

    for r_idx, row in enumerate(df_semanal.itertuples(), start=4):
        ws_as.cell(row=r_idx, column=1, value=row.semana)
        ws_as.cell(row=r_idx, column=2, value=row.avance_planificado_semanal)
        ws_as.cell(row=r_idx, column=3, value=row.avance_real_semanal)
        ws_as.cell(row=r_idx, column=4, value=row.avance_planificado_acumulado)
        ws_as.cell(row=r_idx, column=5, value=row.avance_real_acumulado)
        ws_as.cell(row=r_idx, column=6, value=row.diferencia_acumulada)
        ws_as.cell(row=r_idx, column=7, value=row.estado_avance)

    # 2. Avance_Especialidad
    ws_ae = wb.create_sheet(title="Avance_Especialidad")
    headers_ae = ["Especialidad", "Avance Planificado (%)", "Avance Real (%)", "Diferencia (%)"]
    for c_idx, name in enumerate(headers_ae, start=1):
        ws_ae.cell(row=3, column=c_idx, value=name)
    for r_idx, row in enumerate(df_especialidad.itertuples(), start=4):
        ws_ae.cell(row=r_idx, column=1, value=row.especialidad)
        ws_ae.cell(row=r_idx, column=2, value=row.avance_planificado)
        ws_ae.cell(row=r_idx, column=3, value=row.avance_real)
        ws_ae.cell(row=r_idx, column=4, value=row.diferencia)

    # 3. Partidas_Encargado
    ws_pe = wb.create_sheet(title="Partidas_Encargado")
    headers_pe = ["Encargado", "Cantidad de Partidas"]
    for c_idx, name in enumerate(headers_pe, start=1):
        ws_pe.cell(row=3, column=c_idx, value=name)
    for r_idx, row in enumerate(df_partidas_encargado.itertuples(), start=4):
        ws_pe.cell(row=r_idx, column=1, value=row.encargado)
        ws_pe.cell(row=r_idx, column=2, value=row.cantidad_partidas)

    # 4. Avance_Encargado
    ws_aen = wb.create_sheet(title="Avance_Encargado")
    headers_aen = ["Encargado", "Avance Planificado (%)", "Avance Real (%)", "Diferencia (%)"]
    for c_idx, name in enumerate(headers_aen, start=1):
        ws_aen.cell(row=3, column=c_idx, value=name)
    for r_idx, row in enumerate(df_avance_encargado.itertuples(), start=4):
        ws_aen.cell(row=r_idx, column=1, value=row.encargado)
        ws_aen.cell(row=r_idx, column=2, value=row.avance_planificado)
        ws_aen.cell(row=r_idx, column=3, value=row.avance_real)
        ws_aen.cell(row=r_idx, column=4, value=row.diferencia)

    # 5. Validaciones
    ws_val = wb.create_sheet(title="Validaciones")
    headers_val = ["Validación", "Valor Esperado", "Valor Obtenido", "Resultado"]
    for c_idx, name in enumerate(headers_val, start=1):
        ws_val.cell(row=3, column=c_idx, value=name)

    fila_actual = 4
    for key, vdata in validaciones_dict.items():
        ws_val.cell(row=fila_actual, column=1, value=vdata["descripcion"])
        ws_val.cell(row=fila_actual, column=2, value=vdata["esperado"])
        ws_val.cell(row=fila_actual, column=3, value=vdata["obtenido"])
        ws_val.cell(row=fila_actual, column=4, value="PASÓ" if vdata["estado"] else "FALLÓ")
        fila_actual += 1

    todas_ok = all(v["estado"] for v in validaciones_dict.values())
    ws_val.cell(row=fila_actual + 1, column=1, value="ESTADO GENERAL DE LA BASE DE DATOS:").font = Font(name="Segoe UI", size=10, bold=True)
    ws_val.cell(row=fila_actual + 1, column=4, value="PASÓ" if todas_ok else "FALLÓ")

    # Aplicar estilos
    aplicar_estilo_general_excel(ws_as, "Control de Avance Físico Semanal de la Obra", idx_header=3, tipo_tabla="semanal")
    aplicar_estilo_general_excel(ws_ae, "Avance Acumulado por Especialidad", idx_header=3, tipo_tabla="especialidad")
    aplicar_estilo_general_excel(ws_pe, "Cantidad de Partidas Únicas por Encargado", idx_header=3, tipo_tabla="partidas_encargado")
    aplicar_estilo_general_excel(ws_aen, "Comparativa de Desempeño por Encargado", idx_header=3, tipo_tabla="avance_encargado")
    aplicar_estilo_general_excel(ws_val, "Reporte de Calidad de Datos (10 Validaciones)", idx_header=3, tipo_tabla="validaciones")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def generar_csv_tabla(df_tabla: pd.DataFrame) -> bytes:
    """Genera archivo CSV codificado en UTF-8 con BOM para visualización correcta en Excel."""
    return df_tabla.to_csv(index=False, sep=";", encoding="utf-8-sig").encode("utf-8-sig")


def generar_zip_graficos(dict_bytes_graficos: Dict[str, bytes]) -> bytes:
    """
    Empaqueta múltiples gráficos PNG en un archivo ZIP en memoria.
    """
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for nombre_archivo, bytes_img in dict_bytes_graficos.items():
            zip_file.writestr(nombre_archivo, bytes_img)
    zip_buffer.seek(0)
    return zip_buffer.getvalue()
